import os
import json
import re
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = "courses"
DOCX_DIR = "docx_preview"
ALLOWED_TYPES = {"choice", "matching", "sorting"}


# ================= JSON =================
def load_json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except:
        return {}


def simplify_name(name):
    m = re.match(r"0*(\d+)(.*)", name)
    return f"{int(m.group(1))}{m.group(2)}".strip() if m else name


# ================= STYLES =================
def setup_styles(doc):
    doc.styles["Title"].font.size = Pt(28)

    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(26, 82, 118)

    h2 = doc.styles["Heading 2"]
    h2.font.size = Pt(16)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(17, 122, 101)

    h3 = doc.styles["Heading 3"]
    h3.font.size = Pt(14)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(176, 58, 46)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)


# ================= BLOCKS =================
def add_hint(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("💡 Подсказка: ")
    r.bold = True
    p.add_run(text)

    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F7")
    p._p.get_or_add_pPr().append(shd)


def add_answer_title(doc):
    p = doc.add_paragraph("Ответ:")
    p.runs[0].bold = True
    p.paragraph_format.space_before = Pt(6)


def add_bullet_with_bold_left(doc, text):
    p = doc.add_paragraph(style="List Bullet")

    if "—" in text:
        left, right = text.split("—", 1)
        r1 = p.add_run(left.strip())
        r1.bold = True
        p.add_run(" — " + right.strip())
    else:
        p.add_run(text)


# ================= HTML → DOCX =================
def render_html(doc, html):
    soup = BeautifulSoup(html or "", "html.parser")

    for el in soup.find_all(recursive=False):
        if el.name == "p":
            text = el.get_text(" ", strip=True)
            if text.lower().startswith("подсказка"):
                add_hint(doc, text.replace("Подсказка:", "").strip())
            else:
                doc.add_paragraph(text)

        elif el.name == "ul":
            for li in el.find_all("li"):
                add_bullet_with_bold_left(doc, li.get_text(strip=True))


# ================= TOC =================
def add_table_of_contents(doc):
    """
    Кликабельное оглавление в конце
    """
    doc.add_page_break()


    p = doc.add_paragraph()
    r = p.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)

    doc.add_page_break()


# ================= MAIN =================
def create_docx(course_dir):
    if not os.path.exists(DOCX_DIR):
        os.makedirs(DOCX_DIR)

    course = os.path.basename(course_dir)
    doc = Document()
    setup_styles(doc)

    # --- ТИТУЛЬНАЯ СТРАНИЦА ---
    doc.add_heading(course, 0)
    doc.add_paragraph("Автор: Алексей Павлов").italic = True

    cover = os.path.join(course_dir, "cover.png")
    if os.path.exists(cover):
        doc.add_picture(cover, width=Inches(5))

    doc.add_page_break()

    # --- КОНТЕНТ ---
    modules = sorted(os.listdir(course_dir))
    for module in modules:
        mp = os.path.join(course_dir, module)
        if not os.path.isdir(mp):
            continue

        doc.add_heading(f"📂 Модуль {simplify_name(module)}", 1)

        for lesson in sorted(os.listdir(mp)):
            lp = os.path.join(mp, lesson)
            if not os.path.isdir(lp):
                continue

            doc.add_heading(f"📘 Урок {simplify_name(lesson)}", 2)

            steps = sorted(f for f in os.listdir(lp) if f.endswith(".json"))
            for idx, step_file in enumerate(steps, 1):
                step = load_json(os.path.join(lp, step_file))
                block = step.get("block", {})
                if block.get("name") not in ALLOWED_TYPES:
                    continue

                doc.add_heading(f"💡 Миссия {idx}", 3)
                render_html(doc, block.get("text", ""))

                source = block.get("source", {})
                add_answer_title(doc)

                if block["name"] == "choice":
                    for o in source.get("options", []):
                        if o.get("is_correct"):
                            add_bullet_with_bold_left(doc, o["text"])

                elif block["name"] == "matching":
                    for p in source.get("pairs", []):
                        add_bullet_with_bold_left(
                            doc,
                            f'{p["first"]} — {p["second"]}'
                        )

                elif block["name"] == "sorting":
                    for o in source.get("options", []):
                        add_bullet_with_bold_left(doc, o["text"])

    # --- ОГЛАВЛЕНИЕ В КОНЦЕ ---
    add_table_of_contents(doc)

    out = os.path.join(DOCX_DIR, f"{course}.docx")
    doc.save(out)
    print(f"✅ DOCX создан: {out}")


# ================= RUN =================
if __name__ == "__main__":
    for c in os.listdir(BASE_DIR):
        p = os.path.join(BASE_DIR, c)
        if os.path.isdir(p):
            create_docx(p)
